"""Import the five supplied DOCX articles; keep source text, tables and figures."""
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
import json

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path('/Users/lfbamy/Documents/工作/2026.网站建设/blog/文章/5篇已完成')
SPECS = [
 ('化學發泡與超臨界發泡_原理差異與選擇指南.docx','chemical-foaming-vs-supercritical-foaming','chemical-supercritical','FOAMING TECHNOLOGY / 發泡技術'),
 ('生質含量_碳14_可生物降解_專業解析(1).docx','biobased-carbon14-biodegradability','biobased-carbon14','BIOBASED MATERIALS / 生質材料'),
 ('EVA發泡為什麼會收縮_IUS-4065完整技術分析.docx','eva-foam-shrinkage-ius-4065','eva-shrinkage','DIMENSIONAL STABILITY / 尺寸穩定'),
 ('EVA添加POE與SEBS_雙螺桿分散及發泡配方技術指南.docx','eva-poe-sebs-compounding','eva-poe-sebs','COMPOUNDING / 共混改質'),
 ('IUS-4065_超軟POE低硬度與尺寸穩定技術文章.docx','ius-4065-soft-poe-stability','ius-soft-poe','MATERIAL DEVELOPMENT / 材料開發'),
]
result=[]
for filename, slug, cover, category in SPECS:
 d=Document(SOURCE/filename)
 media=ROOT/'public/images/articles'/slug
 media.mkdir(parents=True,exist_ok=True)
 sections=[{'title':'導讀 / Introduction','text':'','blocks':[]}]
 titles=[p.text.strip() for p in d.paragraphs if p.style.name=='Title' and p.text.strip()]; normal=[]; img_count=0
 for block in d.iter_inner_content():
  if hasattr(block,'text'):
   text=block.text.strip(); style=block.style.name
   if style=='Title':
    continue
   if style=='Heading 1' and text:
    if not sections[-1]['blocks'] and sections[-1]['title']!='導讀 / Introduction':sections[-1]['title']+='\n'+text
    else:sections.append({'title':text,'text':'','blocks':[]})
   elif text:
    kind='heading' if style.startswith('Heading') else 'list' if style.startswith('List') else 'paragraph'
    sections[-1]['blocks'].append({'type':kind,'text':text})
    if style=='Normal':normal.append(text)
   for blip in block._p.xpath('.//a:blip'):
    rid=blip.get(qn('r:embed'))
    if not rid:continue
    part=d.part.related_parts[rid]; ext=str(part.partname).split('.')[-1]
    img_count+=1; name=f'figure-{img_count}.{ext}'; (media/name).write_bytes(part.blob)
    sections[-1]['blocks'].append({'type':'image','src':f'/images/articles/{slug}/{name}','text':f'原文圖 {img_count}｜{titles[0]}'})
  else:
   rows=[]
   for tr in block._tbl.tr_lst:
    row=[]
    for tc in tr.tc_lst:
     paragraphs=[''.join(p.itertext()) for p in []] # text collected once via Word text nodes below
     paragraphs=[''.join(p.xpath('.//w:t/text()')) for p in tc.xpath('./w:p')]
     span=tc.xpath('./w:tcPr/w:gridSpan/@w:val')
     row.append({'text':'\n'.join(paragraphs),'colSpan':int(span[0]) if span else 1})
    rows.append(row)
   sections[-1]['blocks'].append({'type':'table','rows':rows})
 result.append({'slug':slug,'title':titles[0],'englishTitle':titles[1] if len(titles)>1 else '', 'summary':normal[0] if normal else titles[0], 'category':category,'image':f'/images/articles/{cover}-cover.png','alt':f'{titles[0]}｜材料實驗室主題示意圖','sections':sections,'sourceFile':filename,'illustrativeCover':True})
 print(filename, 'sections',len(sections),'tables',len(d.tables),'figures',img_count)
out=ROOT/'app/content/technical-articles.json';out.parent.mkdir(exist_ok=True)
out.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
